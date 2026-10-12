"""Read supplied files without modifying them; retain audit extracts and hashes."""
from pathlib import Path
import json, hashlib, sys, importlib.metadata
import pandas as pd
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs' / 'v6_result_to_draw_20261011'
INPUT = Path(r'C:\Users\SportEdge\Downloads\drive-download-20261011T181047Z-1-001')

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT/'inputs').mkdir(exist_ok=True)
    manifest=[]
    for p in INPUT.iterdir():
        if p.suffix.lower() not in ['.docx','.xlsx','.xlsm']: continue
        manifest.append({'path':str(p),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
        if p.suffix=='.docx':
            doc=Document(p)
            lines=[q.text for q in doc.paragraphs if q.text]
            for t in doc.tables:
                lines.extend(' | '.join(c.text for c in r.cells) for r in t.rows)
            (OUT/'inputs'/f'{p.stem}.txt').write_text('\n'.join(lines),encoding='utf8')
        else:
            sheets=pd.read_excel(p,sheet_name=None,header=None,dtype=object)
            data={k:v.where(v.notna(),None).values.tolist() for k,v in sheets.items()}
            (OUT/'inputs'/f'{p.stem}.json').write_text(json.dumps(data,ensure_ascii=False,default=str),encoding='utf8')
            print(p.name, {k:v.shape for k,v in sheets.items()})
            if p.suffix=='.xlsx': print(json.dumps(data,ensure_ascii=False,default=str)[:8500])
    csv=ROOT/'outputs/etiquetado_result_to_draw_20261011/Match_centre_base_modelos_V2_etiquetada_result_to_draw.csv'
    d=pd.read_csv(csv,low_memory=False)
    manifest.append({'path':str(csv),'bytes':csv.stat().st_size,'sha256':hashlib.sha256(csv.read_bytes()).hexdigest()})
    (OUT/'inputs/manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf8')
    print('DATA',d.shape, 'ids',d.id.nunique(),'target',d.result_to_draw.value_counts(dropna=False).to_dict())
    print('COLS',d.columns.tolist())
    print('DATES',d.starting_at.min(),d.starting_at.max())
    for package in ['numpy','pandas','scikit-learn','xgboost','catboost','lightgbm','fastapi','uvicorn','joblib','scipy','optuna']:
        try: print(package,importlib.metadata.version(package))
        except: print(package,'unavailable')
    a=json.loads((ROOT/'audit/labeling_validation.json').read_text(encoding='utf8'))
    print('LABEL AUDIT', {k:v for k,v in a.items() if k!='group_details'})

if __name__=='__main__': main()
